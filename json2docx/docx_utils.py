from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from replacements import normalize_replacement

PLACEHOLDER_RE = re.compile(r"<<[^>]+>>|{{[^}]+}}")


def read_docx_paragraph_text(docx_path: Path) -> str:
    """Return only BODY paragraph text. No tables/headers/footers."""
    doc = Document(str(docx_path))
    return "\n\n".join(p.text.rstrip() for p in doc.paragraphs).strip()


def extract_placeholders_from_docx(docx_path: Path) -> List[str]:
    """Return unique placeholders found in body paragraphs, in document order."""
    seen: set[str] = set()
    ordered: List[str] = []
    for match in PLACEHOLDER_RE.finditer(read_docx_paragraph_text(docx_path)):
        token = match.group(0)
        if token not in seen:
            seen.add(token)
            ordered.append(token)
    return ordered


def _set_xml_space_preserve(t_elem: OxmlElement, text: str) -> None:
    if text.startswith(" ") or text.endswith(" "):
        t_elem.set(qn("xml:space"), "preserve")


def _insert_run_after(after_run: Run, text: str, rpr: Optional[OxmlElement]) -> Run:
    new_r = OxmlElement("w:r")
    if rpr is not None:
        new_r.append(deepcopy(rpr))
    new_t = OxmlElement("w:t")
    _set_xml_space_preserve(new_t, text)
    new_t.text = text
    new_r.append(new_t)
    after_run._r.addnext(new_r)
    return Run(new_r, after_run._parent)


def _collect_char_map(runs: List[Run]) -> List[Tuple[int, int]]:
    return [(ri, off) for ri, r in enumerate(runs) for off in range(len(r.text))]


def _mark_bold_regions(text: str, phrases: Iterable[str]) -> List[bool]:
    mask = [False] * len(text)
    seen: set[str] = set()
    for phrase in phrases:
        if not phrase or phrase in seen:
            continue
        seen.add(phrase)
        start = 0
        while True:
            idx = text.find(phrase, start)
            if idx == -1:
                break
            for j in range(idx, min(idx + len(phrase), len(mask))):
                mask[j] = True
            start = idx + 1
    return mask


def _split_by_bold_mask(text: str, mask: List[bool]) -> List[Tuple[str, bool]]:
    if not text:
        return []
    if len(mask) != len(text):
        mask = [False] * len(text)
    parts: List[Tuple[str, bool]] = []
    cur = [text[0]]
    cur_bold = mask[0]
    for i in range(1, len(text)):
        if mask[i] == cur_bold:
            cur.append(text[i])
        else:
            parts.append(("".join(cur), cur_bold))
            cur = [text[i]]
            cur_bold = mask[i]
    parts.append(("".join(cur), cur_bold))
    return parts


def _clear_paragraph_runs(paragraph: Paragraph) -> Optional[OxmlElement]:
    runs = list(paragraph.runs)
    rpr = runs[0]._r.rPr if runs else None
    for run in runs:
        run._element.getparent().remove(run._element)
    return rpr


def _append_text_runs(
    paragraph: Paragraph,
    text: str,
    bold_phrases: List[str],
    rpr: Optional[OxmlElement],
    *,
    force_bold: bool = False,
) -> None:
    text = text.replace("\r", "").replace("\n", " ").strip()
    if not text:
        return
    anchor_run: Run | None = None
    mask = _mark_bold_regions(text, bold_phrases)
    for seg_text, seg_bold in _split_by_bold_mask(text, mask):
        if anchor_run is None:
            new_r = OxmlElement("w:r")
            if rpr is not None:
                new_r.append(deepcopy(rpr))
            new_t = OxmlElement("w:t")
            _set_xml_space_preserve(new_t, seg_text)
            new_t.text = seg_text
            new_r.append(new_t)
            paragraph._p.append(new_r)
            run = Run(new_r, paragraph)
            run.bold = force_bold or bool(seg_bold)
            anchor_run = run
        else:
            new_run = _insert_run_after(anchor_run, seg_text, rpr)
            new_run.bold = force_bold or bool(seg_bold)
            anchor_run = new_run


def _insert_paragraph_after(reference: Paragraph, text: str, bold_phrases: List[str]) -> Paragraph:
    new_p = OxmlElement("w:p")
    reference._p.addnext(new_p)
    new_para = Paragraph(new_p, reference._parent)
    if reference.style is not None:
        new_para.style = reference.style
    rpr = _clear_paragraph_runs(new_para)
    _append_text_runs(new_para, text, bold_phrases, rpr)
    return new_para


def _insert_parts_after_paragraph(
    paragraph: Paragraph,
    parts: List[str],
    bold_phrases: List[str],
) -> None:
    if not parts:
        return
    runs = list(paragraph.runs)
    force_bold = any(r.bold is True for r in runs)
    rpr = runs[0]._r.rPr if runs else None
    _clear_paragraph_runs(paragraph)
    _append_text_runs(paragraph, parts[0], bold_phrases, rpr, force_bold=force_bold)
    anchor = paragraph
    for part in parts[1:]:
        anchor = _insert_paragraph_after(anchor, part, bold_phrases)


def _replace_multiline_placeholder(
    paragraph: Paragraph,
    placeholder: str,
    replacement: str,
    bold_phrases: List[str],
) -> bool:
    full_text = "".join(r.text for r in paragraph.runs)
    if placeholder not in full_text:
        return False

    parts = [p.strip() for p in replacement.replace("\r", "").split("\n\n") if p.strip()]
    if not parts:
        return _replace_once_in_paragraph(paragraph, placeholder, "", bold_phrases)

    before, _, after = full_text.partition(placeholder)
    if before or after:
        if not _replace_once_in_paragraph(paragraph, placeholder, parts[0], bold_phrases):
            return False
        anchor = paragraph
        for part in parts[1:]:
            anchor = _insert_paragraph_after(anchor, part, bold_phrases)
        return True

    _insert_parts_after_paragraph(paragraph, parts, bold_phrases)
    return True


def _replace_once_in_paragraph(
    paragraph: Paragraph, placeholder: str, replacement: str, bold_phrases: List[str]
) -> bool:
    if "\n\n" in (replacement or ""):
        return _replace_multiline_placeholder(paragraph, placeholder, replacement, bold_phrases)

    runs = list(paragraph.runs)
    if not runs:
        return False

    full_text = "".join(r.text for r in runs)
    idx = full_text.find(placeholder)
    if idx == -1:
        return False

    replacement = replacement.replace("\r", " ").replace("\n", " ")
    mapping = _collect_char_map(runs)
    if idx + len(placeholder) > len(mapping):
        paragraph.text = full_text.replace(placeholder, replacement, 1)
        return True

    start_ri, start_off = mapping[idx]
    end_ri, end_off = mapping[idx + len(placeholder) - 1]
    force_bold = any(runs[ri].bold is True for ri in range(start_ri, end_ri + 1))

    start_run = runs[start_ri]
    end_run = runs[end_ri]
    rpr = start_run._r.rPr
    before = start_run.text[:start_off]
    after = end_run.text[end_off + 1 :]

    start_run.text = before
    for ri in range(start_ri + 1, end_ri + 1):
        runs[ri].text = ""
    if start_ri != end_ri:
        end_run.text = after

    anchor = start_run
    mask = _mark_bold_regions(replacement, bold_phrases)
    for seg_text, seg_bold in _split_by_bold_mask(replacement, mask):
        new_run = _insert_run_after(anchor, seg_text, rpr)
        new_run.bold = force_bold or bool(seg_bold)
        anchor = new_run

    if start_ri == end_ri and after:
        after_run = _insert_run_after(anchor, after, rpr)
        after_run.bold = start_run.bold

    return True


def replace_placeholders_in_docx(template_docx: Path, replacements: Dict[str, Dict], out_docx: Path) -> None:
    """Replace all placeholder occurrences in BODY paragraphs only."""
    doc = Document(str(template_docx))
    normalized = {k: normalize_replacement(v) for k, v in replacements.items()}
    keys = sorted(normalized, key=len, reverse=True)

    for p in doc.paragraphs:
        if "<<" not in p.text and "{{" not in p.text:
            continue
        changed = True
        while changed:
            changed = False
            full = "".join(r.text for r in p.runs)
            for key in keys:
                if key in full:
                    val = normalized[key]
                    if _replace_once_in_paragraph(
                        p, key, val["text"], list(val.get("bold-words", []))
                    ):
                        changed = True
                        break

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))
